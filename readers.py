import io
from pathlib import Path
import docx
import ebooklib
import pandas as pd
from bs4 import BeautifulSoup
from ebooklib import epub
from pptx import Presentation
from pypdf import PdfReader

def read_pdf(data):
    for i, page in enumerate(PdfReader(io.BytesIO(data)).pages, start=1):
        yield f"Page {i}", page.extract_text() or ""

def read_docx(data):
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        parts += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
    yield "Document", "\n".join(parts)

def read_pptx(data):
    for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                parts += [" | ".join(cell.text for cell in row.cells) for row in shape.table.rows]
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
        yield f"Slide {i}", "\n".join(parts)

def rows_to_text(df):
    # "column: value" per row keeps the headers attached to every chunk.
    return "\n".join(
        ", ".join(f"{key}: {value}" for key, value in row.items() if pd.notna(value))
        for row in df.to_dict("records")
    )

def read_excel(data):
    for sheet, df in pd.read_excel(io.BytesIO(data), sheet_name=None).items():
        yield f"Sheet {sheet}", rows_to_text(df)

def read_csv(data, sep=","):
    yield "Table", rows_to_text(pd.read_csv(io.BytesIO(data), sep=sep))

def html_to_text(html):
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav"]):
        tag.decompose()
    heading = soup.find(["h1", "h2", "h3"])
    title = heading.get_text(" ", strip=True) if heading else None
    return title, soup.get_text(" ", strip=True)

def read_html(data):
    title, text = html_to_text(data)
    yield title or "Page", text

def read_epub(data):
    book = epub.read_epub(io.BytesIO(data))
    for i, item in enumerate(book.get_items_of_type(ebooklib.ITEM_DOCUMENT), start=1):
        title, text = html_to_text(item.get_content())
        yield title or f"Section {i}", text

def read_text(data):
    yield "Text", data.decode("utf-8-sig", errors="replace")

READERS = {
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".pptx": read_pptx,
    ".xlsx": read_excel,
    ".xlsm": read_excel,
    ".xls": read_excel,
    ".csv": read_csv,
    ".tsv": lambda data: read_csv(data, sep="\t"),
    ".html": read_html,
    ".htm": read_html,
    ".epub": read_epub,
    ".txt": read_text,
    ".md": read_text,
    ".json": read_text,
    ".xml": read_text,
}

def read_file(name, data):
    reader = READERS[Path(name).suffix.lower()]
    return [(location, text) for location, text in reader(data) if text.strip()]
